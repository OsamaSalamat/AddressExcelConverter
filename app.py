from __future__ import annotations

import csv
import re
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT_COLUMNS = (
    "Customer_Name",
    "AddressOne",
    "Country",
    "Postal",
    "City",
    "State",
    "Service",
    "Package_Type",
    "Billing",
    "Package_Weight",
)
FIXED_COLUMNS = ("Service", "Package_Type", "Billing", "Package_Weight")
DELIMITERS = (",", "\t", ";", "|")
ADDRESS_HEADER = ("name", "street", "city", "state", "zip")


class InputError(ValueError):
    """Raised when an input file cannot be mapped to the required fields."""


def _delimiter_for(line: str) -> str | None:
    counts = {delimiter: line.count(delimiter) for delimiter in DELIMITERS}
    delimiter, count = max(counts.items(), key=lambda item: item[1])
    return delimiter if count else None


def parse_addresses(text: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    errors: list[str] = []

    for line_number, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            continue

        delimiter = _delimiter_for(line)
        if delimiter is None:
            errors.append(f"Line {line_number}: no comma, tab, semicolon, or pipe delimiter found.")
            continue

        try:
            fields = next(csv.reader([line], delimiter=delimiter, skipinitialspace=True))
        except csv.Error as error:
            errors.append(f"Line {line_number}: {error}")
            continue

        fields = [field.strip() for field in fields]
        normalized = tuple(re.sub(r"[^a-z]", "", field.lower()) for field in fields)
        if normalized == ADDRESS_HEADER:
            continue

        if len(fields) < 5:
            errors.append(f"Line {line_number}: expected Name, Street, City, State, Zip; found {len(fields)} fields.")
            continue

        # Extra delimiter-separated parts are treated as part of the street.
        name = fields[0]
        street = delimiter.join(fields[1:-3]).strip()
        city, state, postal = fields[-3:]
        if not all((name, street, city, state, postal)):
            errors.append(f"Line {line_number}: one or more required address fields are blank.")
            continue

        rows.append(
            {
                "Customer_Name": name,
                "AddressOne": street,
                "Country": "US",
                "Postal": postal,
                "City": city,
                "State": state,
            }
        )

    if errors:
        preview = "\n".join(errors[:20])
        if len(errors) > 20:
            preview += f"\n... and {len(errors) - 20} more issue(s)."
        raise InputError(preview)
    if not rows:
        raise InputError("No address rows were found in the text file.")
    return rows


def read_fixed_fields(csv_path: Path) -> dict[str, str]:
    try:
        with csv_path.open("r", encoding="utf-8-sig", newline="") as file:
            content = [row for row in csv.reader(file) if any(value.strip() for value in row)]
    except (OSError, csv.Error) as error:
        raise InputError(f"Could not read the fixed-fields CSV: {error}") from error

    if not content:
        raise InputError("The fixed-fields CSV is empty.")

    first = [value.strip() for value in content[0]]
    normalized_headers = tuple(re.sub(r"[^a-z]", "", value.lower()) for value in first)
    expected_headers = tuple(re.sub(r"[^a-z]", "", value.lower()) for value in FIXED_COLUMNS)
    has_header = normalized_headers == expected_headers
    values_rows = content[1:] if has_header else content

    if len(values_rows) != 1:
        raise InputError("The fixed-fields CSV must contain exactly one data row, optionally preceded by a header row.")
    values = [value.strip() for value in values_rows[0]]
    if len(values) != len(FIXED_COLUMNS) or any(not value for value in values):
        raise InputError(
            "The fixed-fields CSV must have four nonblank values in this order: "
            "Service, Package_Type, Billing, Package_Weight."
        )
    return dict(zip(FIXED_COLUMNS, values))


def write_workbook(rows: list[dict[str, str]], fixed_fields: dict[str, str], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Addresses"
    sheet.append(OUTPUT_COLUMNS)

    for address in rows:
        values = {**address, **fixed_fields}
        sheet.append([values[column] for column in OUTPUT_COLUMNS])

    header_fill = PatternFill("solid", fgColor="235B4E")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index in range(1, len(OUTPUT_COLUMNS) + 1):
        values = [str(sheet.cell(row, column_index).value or "") for row in range(1, sheet.max_row + 1)]
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(map(len, values)) + 2, 36)

    workbook.save(output_path)


class AddressExcelApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Address Excel Converter")
        self.root.geometry("760x440")
        self.root.minsize(620, 380)
        self.address_path = tk.StringVar()
        self.fixed_path = tk.StringVar()
        self.status = tk.StringVar(value="Choose the address text file and fixed-fields CSV to begin.")
        self._build_ui()

    def _build_ui(self) -> None:
        frame = ttk.Frame(self.root, padding=22)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)

        ttk.Label(frame, text="Address Excel Converter", font=("Segoe UI", 18, "bold")).grid(
            row=0, column=0, columnspan=3, sticky="w", pady=(0, 6)
        )
        ttk.Label(
            frame,
            text="Convert one delimited address per line into the client’s Excel column layout.",
        ).grid(row=1, column=0, columnspan=3, sticky="w", pady=(0, 22))

        ttk.Label(frame, text="Address text file").grid(row=2, column=0, sticky="w", padx=(0, 12), pady=8)
        ttk.Entry(frame, textvariable=self.address_path).grid(row=2, column=1, sticky="ew", pady=8)
        ttk.Button(frame, text="Browse...", command=self._choose_address).grid(row=2, column=2, padx=(10, 0), pady=8)

        ttk.Label(frame, text="Fixed-fields CSV").grid(row=3, column=0, sticky="w", padx=(0, 12), pady=8)
        ttk.Entry(frame, textvariable=self.fixed_path).grid(row=3, column=1, sticky="ew", pady=8)
        ttk.Button(frame, text="Browse...", command=self._choose_fixed).grid(row=3, column=2, padx=(10, 0), pady=8)

        ttk.Separator(frame).grid(row=4, column=0, columnspan=3, sticky="ew", pady=18)
        ttk.Button(frame, text="Process and create Excel", command=self._process).grid(
            row=5, column=0, columnspan=3, sticky="ew", ipady=7
        )

        ttk.Label(frame, textvariable=self.status, wraplength=700).grid(
            row=6, column=0, columnspan=3, sticky="nw", pady=(18, 8)
        )
        ttk.Label(
            frame,
            text="Text rows: Name, Street, City, State, Zip. Supported delimiters: comma, tab, semicolon, pipe.",
            wraplength=700,
        ).grid(row=7, column=0, columnspan=3, sticky="w", pady=(10, 0))

    def _choose_address(self) -> None:
        path = filedialog.askopenfilename(
            title="Choose address text file", filetypes=(("Text files", "*.txt"), ("All files", "*.*"))
        )
        if path:
            self.address_path.set(path)

    def _choose_fixed(self) -> None:
        path = filedialog.askopenfilename(
            title="Choose fixed-fields CSV", filetypes=(("CSV files", "*.csv"), ("All files", "*.*"))
        )
        if path:
            self.fixed_path.set(path)

    def _process(self) -> None:
        address_file = Path(self.address_path.get())
        fixed_file = Path(self.fixed_path.get())
        if not address_file.is_file() or not fixed_file.is_file():
            messagebox.showerror("Missing input", "Choose a valid address text file and fixed-fields CSV.")
            return

        output_name = f"{address_file.stem}_converted.xlsx"
        output = filedialog.asksaveasfilename(
            title="Save Excel workbook",
            initialfile=output_name,
            defaultextension=".xlsx",
            filetypes=(("Excel workbook", "*.xlsx"),),
        )
        if not output:
            return

        try:
            text = address_file.read_text(encoding="utf-8-sig")
            addresses = parse_addresses(text)
            fixed_fields = read_fixed_fields(fixed_file)
            write_workbook(addresses, fixed_fields, Path(output))
        except (OSError, InputError, PermissionError) as error:
            self.status.set("Processing failed. Correct the input and try again.")
            messagebox.showerror("Could not create workbook", str(error))
            return
        except Exception as error:
            self.status.set("Processing failed. Correct the input and try again.")
            messagebox.showerror("Could not create workbook", f"Unexpected error: {error}")
            return

        self.status.set(f"Created {output} with {len(addresses):,} address row(s).")
        messagebox.showinfo("Finished", f"Created Excel workbook with {len(addresses):,} address row(s).")


def main() -> None:
    root = tk.Tk()
    AddressExcelApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()