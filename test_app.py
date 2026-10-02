import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app import FIXED_COLUMNS, OUTPUT_COLUMNS, InputError, parse_addresses, read_fixed_fields, write_workbook


class AddressConverterTests(unittest.TestCase):
    def test_parse_supported_delimiters_and_comma_in_street(self) -> None:
        rows = parse_addresses(
            "Name,Street,City,State,Zip\n"
            'Jane Doe,"12 Main St, Apt 4",Springfield,IL,62704\n'
            "John Doe|5 Oak Road|Madison|WI|53703"
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["AddressOne"], "12 Main St, Apt 4")
        self.assertEqual(rows[0]["Postal"], "62704")
        self.assertEqual(rows[1]["Country"], "US")

    def test_invalid_address_line_is_reported(self) -> None:
        with self.assertRaisesRegex(InputError, "Line 1"):
            parse_addresses("Jane Doe, 12 Main St, Springfield")

    def test_fixed_fields_and_workbook_column_order(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            csv_path = Path(directory) / "fixed.csv"
            output_path = Path(directory) / "output.xlsx"
            csv_path.write_text(
                "Service,Package_Type,Billing,Package_Weight\nGround,Parcel,Prepaid,1\n",
                encoding="utf-8",
            )
            fixed = read_fixed_fields(csv_path)
            addresses = parse_addresses("Jane Doe,12 Main St,Springfield,IL,62704")
            write_workbook(addresses, fixed, output_path)

            workbook = load_workbook(output_path, read_only=True)
            try:
                sheet = workbook.active
                self.assertEqual(tuple(cell.value for cell in next(sheet.iter_rows())), OUTPUT_COLUMNS)
                self.assertEqual(tuple(fixed), FIXED_COLUMNS)
                self.assertEqual(sheet.max_row, 2)
            finally:
                workbook.close()

    def test_handles_one_thousand_rows(self) -> None:
        text = "\n".join(
            f"Customer {index}, {index} Main St, Springfield, IL, {60000 + index}"
            for index in range(1000)
        )
        addresses = parse_addresses(text)
        fixed = dict(zip(FIXED_COLUMNS, ("Ground", "Parcel", "Prepaid", "1")))
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "thousand.xlsx"
            write_workbook(addresses, fixed, output_path)
            workbook = load_workbook(output_path, read_only=True)
            try:
                sheet = workbook.active
                self.assertEqual(sheet.max_row, 1001)
                self.assertEqual(sheet.max_column, len(OUTPUT_COLUMNS))
                self.assertEqual(tuple(cell.value for cell in next(sheet.iter_rows(min_row=2))),
                                 ("Customer 0", "0 Main St", "US", "60000", "Springfield", "IL",
                                  "Ground", "Parcel", "Prepaid", "1"))
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()